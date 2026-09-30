"""Bank statement preview: CSV, XLSX and bank HTML exports named XLS.

Cash Trail (IkerJansa44/cash-trail) informed the occurrence-aware import design.
No bank connection, AI call or database write is performed while parsing.
"""
import csv
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from html.parser import HTMLParser
from pathlib import Path
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook


def clean(value):
    return str(value if value is not None else '').strip()


def normalized(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', clean(value).lower()) if not unicodedata.combining(c)).strip()


def cents(value, decimal='auto'):
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = clean(value).replace('\u00a0', '').replace(' ', '').replace('€', '').replace('EUR', '')
        if not raw:
            raise ValueError('Amount is empty')
        if raw.startswith('(') and raw.endswith(')'):
            raw = '-' + raw[1:-1]
        if decimal == ',':
            raw = raw.replace('.', '').replace(',', '.')
        elif decimal == '.':
            raw = raw.replace(',', '')
        elif ',' in raw:
            if '.' in raw and raw.rfind('.') > raw.rfind(','):
                raw = raw.replace(',', '')
            else:
                raw = raw.replace('.', '').replace(',', '.')
    try:
        result = Decimal(raw)
        if not result.is_finite() or abs(result) > Decimal('9999999999'):
            raise ValueError('Amount is outside the supported range')
        return int((result * 100).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except InvalidOperation as error:
        raise ValueError(f'Invalid amount: {clean(value)[:40]}') from error


def parse_date(value, order='day-first'):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = clean(value)
    if re.fullmatch(r'\d+(\.0+)?', text) and 20000 < float(text) < 100000:
        return (date(1899, 12, 30) + timedelta(days=int(float(text)))).isoformat()
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        pass
    formats = ['%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y', '%d/%m/%y'] if order == 'day-first' else ['%m/%d/%Y', '%m-%d-%Y', '%m/%d/%y']
    for fmt in formats:
        try:
            return datetime.strptime(text.split(' ')[0], fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f'Invalid date: {text[:40]}')


class HtmlRows(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == 'tr':
            self.row = []
        elif tag in ('td', 'th') and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(' '.join(''.join(self.cell).split()))
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def decode(content):
    for encoding in ('utf-8-sig', 'utf-16', 'cp1252'):
        try:
            return content.decode(encoding)
        except UnicodeError:
            pass
    raise ValueError('Unable to read the text encoding')


ALIASES = {
    'date': ['date', 'fecha', 'fecha operacion', "data d'operacio", 'data operacio', 'booking date', 'transaction date', 'fecha movimiento', 'completed date', 'started date', 'posting date'],
    'description': ['description', 'descripcion', 'concepto', 'concepte', 'concept', 'details', 'transaction description', 'movimiento'],
    'amount': ['amount', 'importe', 'import', 'cantidad', 'transaction amount'],
    'debit': ['debit', 'cargo', 'cargos', 'debe', 'withdrawal', 'withdrawals'],
    'credit': ['credit', 'abono', 'abonos', 'haber', 'deposit', 'deposits'],
    'balance': ['balance', 'saldo'],
    'currency': ['currency', 'moneda', 'divisa'],
}


def read_table(filename, content, sheet=''):
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.csv', '.xlsx', '.xls'):
        raise ValueError('Choose an Excel (.xlsx), CSV, or bank HTML .xls export. PDF is not supported.')
    if suffix in ('.xls', '.xlsx') and content[:2] == b'PK':
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 100_000_000:
                    raise ValueError('This workbook is too large when expanded')
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            names = book.sheetnames
            if sheet and sheet not in names:
                raise ValueError('Select a worksheet from this file')
            worksheet = book[sheet] if sheet else book.active
            rows = []
            for row in worksheet.iter_rows(values_only=True):
                if len(rows) >= 20000 or len(row) > 150:
                    raise ValueError('Limit: 20,000 rows and 150 columns per import')
                rows.append(list(row))
            chosen = worksheet.title
            book.close()
            return rows, names, chosen
        except (BadZipFile, KeyError) as error:
            raise ValueError('The workbook could not be read') from error
    if suffix == '.xlsx':
        raise ValueError('This file is not a valid .xlsx workbook')
    if content.startswith(b'\xd0\xcf\x11\xe0'):
        raise ValueError('This is an older binary Excel file. Save it as .xlsx or CSV in Excel, then import it.')
    text = decode(content)
    if '<table' in text[:10000].lower() or '<html' in text[:1000].lower():
        parser = HtmlRows()
        parser.feed(text)
        rows = parser.rows
    elif suffix == '.csv':
        try:
            dialect = csv.Sniffer().sniff(text[:12000], delimiters=',;\t|')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ValueError('This .xls export is not an HTML table. Save it as .xlsx or CSV.')
    if len(rows) > 20000 or any(len(r) > 150 for r in rows):
        raise ValueError('Limit: 20,000 rows and 150 columns per import')
    return rows, [], ''


def infer_mapping(headers):
    names = [normalized(x) for x in headers]
    return {field: next((i for i, name in enumerate(names) if name in aliases), -1) for field, aliases in ALIASES.items()}


def preview(filename, content, options):
    rows, sheets, sheet = read_table(filename, content, options.get('sheet', ''))
    if not rows:
        raise ValueError('This file is empty')
    header = options.get('header')
    if header is None:
        header = max(range(min(40, len(rows))), key=lambda i: sum(v >= 0 for v in infer_mapping(rows[i]).values()))
    header = int(header)
    if not 0 <= header < len(rows):
        raise ValueError('Header row is outside this file')
    headers = [clean(x) or f'Column {i + 1}' for i, x in enumerate(rows[header])]
    mapping = options.get('mapping') or infer_mapping(headers)
    mapping = {key: int(mapping.get(key, -1)) for key in ALIASES}
    if any(i < -1 or i >= len(headers) for i in mapping.values()):
        raise ValueError('Column selection is outside the table')
    result = {'headers': headers, 'header': header, 'mapping': mapping, 'sheets': sheets, 'sheet': sheet,
              'sample': [[clean(x) for x in row] for row in rows[header + 1:header + 6]], 'rows': [], 'errors': [], 'filename': Path(filename).name}
    if mapping['date'] < 0 or mapping['description'] < 0 or (mapping['amount'] < 0 and (mapping['debit'] < 0 or mapping['credit'] < 0)):
        result['mapping_required'] = True
        return result
    def cell(row, field):
        i = mapping[field]
        return row[i] if 0 <= i < len(row) else ''
    occurrences = defaultdict(int)
    decimal = options.get('decimal', 'auto')
    for index, row in enumerate(rows[header + 1:], header + 2):
        if not any(clean(x) for x in row):
            continue
        try:
            day = parse_date(cell(row, 'date'), options.get('date_order', 'day-first'))
            description = clean(cell(row, 'description'))
            if not description:
                raise ValueError('Description is empty')
            currency = clean(cell(row, 'currency')).upper() or options.get('currency', 'EUR')
            currency = 'EUR' if currency == '€' else currency
            if currency not in ('EUR', 'AED', 'SCR', 'USD', 'GBP', 'CHF', 'SEK', 'NOK', 'DKK', 'PLN', 'CAD', 'AUD', 'SGD'):
                raise ValueError(f'Unsupported currency: {currency}. Choose a supported account currency.')
            if mapping['amount'] >= 0:
                amount = cents(cell(row, 'amount'), decimal)
            else:
                amount = abs(cents(cell(row, 'credit'), decimal)) if clean(cell(row, 'credit')) else 0
                amount -= abs(cents(cell(row, 'debit'), decimal)) if clean(cell(row, 'debit')) else 0
            if options.get('invert'):
                amount = -amount
            balance = cents(cell(row, 'balance'), decimal) if clean(cell(row, 'balance')) else None
            raw = '|'.join([day, normalized(description), str(amount), str(balance), currency])
            occurrences[raw] += 1
            fingerprint = hashlib.sha256((raw + '|' + str(occurrences[raw])).encode()).hexdigest()
            result['rows'].append({'date': day, 'description': description[:500], 'cents': amount, 'currency': currency, 'balance': balance,
                                   'fingerprint': fingerprint, 'row_number': index})
        except ValueError as error:
            result['errors'].append({'row': index, 'error': str(error)})
    if result['rows']:
        dates = [r['date'] for r in result['rows']]
        result['start'], result['end'] = min(dates), max(dates)
    return result

